from openpyxl import load_workbook
import os

excel_path = os.path.join(os.path.dirname(__file__), '..', 'HIQ_Rate contract.xlsx')
wb = load_workbook(excel_path)
ws = wb.active

print("=" * 80)
print("DETAILED EXCEL STRUCTURE ANALYSIS")
print("=" * 80)

# Read row 18 which seems to be the main header
print("\nRow 18 (Main Header):")
print("-" * 80)
header_row = [cell.value for cell in ws[18]]
for idx, val in enumerate(header_row, 1):
    if val:
        print(f"Column {idx}: {val}")

# Read rows 19-25 to understand the structure
print("\n\nRows 19-25 (Specification rows):")
print("-" * 80)
for row_idx in range(19, 26):
    row = [cell.value for cell in ws[row_idx]]
    non_empty = [(idx, val) for idx, val in enumerate(row, 1) if val]
    if non_empty:
        print(f"\nRow {row_idx}:")
        for col_idx, val in non_empty[:10]:  # Show first 10 non-empty columns
            print(f"  Col {col_idx}: {val}")

# Read sample data rows (starting from row 22)
print("\n\nSample Data Rows (22-35):")
print("-" * 80)
for row_idx in range(22, min(36, ws.max_row + 1)):
    row = [cell.value for cell in ws[row_idx]]
    # Get all non-empty values
    row_data = [str(val) if val is not None else '' for val in row]
    row_str = ' | '.join([val[:25] for val in row_data if val.strip()])
    if row_str.strip():
        print(f"Row {row_idx}: {row_str[:200]}")  # Limit output length

# Try to understand the rate structure
print("\n\nAnalyzing Rate Structure:")
print("-" * 80)
print("Looking for rate values (numeric) in rows 22-50...")
for row_idx in range(22, min(51, ws.max_row + 1)):
    row = [cell.value for cell in ws[row_idx]]
    numeric_values = [(idx, val) for idx, val in enumerate(row, 1) if isinstance(val, (int, float))]
    if numeric_values:
        print(f"\nRow {row_idx} has numeric values:")
        for col_idx, val in numeric_values[:5]:  # Show first 5
            header_val = header_row[col_idx - 1] if col_idx <= len(header_row) else None
            print(f"  Col {col_idx} ({header_val}): {val}")

wb.close()
