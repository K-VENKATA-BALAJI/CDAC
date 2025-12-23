"""
Import Vendor Rate Contract Excel data into MySQL database
Generic version - works for any vendor
"""
from openpyxl import load_workbook
import mysql.connector
from mysql.connector import Error
import os
import sys
import re

# Database configuration
DB_CONFIG = {
    'host': 'localhost',
    'database': 'pcb_rate_contract',
    'user': 'root',
    'password': 'Balu@6688'  # Update with your MySQL password
}

# ============================================
# CONFIGURATION - UPDATE THESE VALUES
# ============================================
VENDOR_NAME = 'VENDOR2'  # Short name/identifier (e.g., "VENDOR2", "ABC_PCB")
EXCEL_FILE_PATH = 'Vendor2_Rate contract.xlsx'  # Path to Excel file

def get_db_connection():
    """Create and return database connection"""
    try:
        connection = mysql.connector.connect(**DB_CONFIG)
        return connection
    except Error as e:
        print(f"Error connecting to MySQL: {e}")
        sys.exit(1)

def parse_excel_data(excel_path, vendor_name):
    """Parse Excel file and extract rate data"""
    wb = load_workbook(excel_path)
    ws = wb.active
    
    # Extract vendor information (rows 6-8)
    vendor_info = {
        'name': vendor_name,
        'company_name': None,
        'address': None,
        'phone': None
    }
    
    for row_idx in range(6, 9):
        row = [cell.value for cell in ws[row_idx]]
        if row[0] == 'Name of the Company' and len(row) > 2:
            vendor_info['company_name'] = str(row[2]) if row[2] else None
        elif row[0] == 'Address' and len(row) > 2:
            vendor_info['address'] = str(row[2]) if row[2] else None
        elif 'Phone' in str(row[0]) and len(row) > 2:
            vendor_info['phone'] = str(row[2]) if row[2] else None
    
    # Extract basic PCB specifications (rows 11-15)
    basic_specs = {}
    for row_idx in range(11, 16):
        row = [cell.value for cell in ws[row_idx]]
        if row[0] == 'PCB Material' and len(row) > 2:
            basic_specs['material'] = str(row[2]) if row[2] else None
        elif row[0] == 'PCB Finish' and len(row) > 2:
            basic_specs['surface_finish'] = str(row[2]) if row[2] else None
        elif row[0] == 'Copper Thickness' and len(row) > 2:
            basic_specs['copper_thickness'] = str(row[2]) if row[2] else None
        elif row[0] == 'Minimum Quantity' and len(row) > 2:
            basic_specs['min_quantity'] = str(row[2]) if row[2] else None
        elif row[0] == 'BBT Charges' and len(row) > 2:
            basic_specs['bbt_charges'] = str(row[2]) if row[2] else None
    
    # Extract track/spacing and via values (rows 19-20)
    row19 = [cell.value for cell in ws[19]]  # Track/Spacing
    row20 = [cell.value for cell in ws[20]]  # Via drill/pad
    
    # Map column indices to track/spacing and via values
    column_specs = {}
    
    # Find track/spacing values
    track_spacing_map = {}
    for col_idx in range(4, len(row19)):
        if row19[col_idx]:
            track_val = str(row19[col_idx]).strip()
            if 'mil' in track_val.lower():
                track_spacing_map[col_idx + 1] = track_val
    
    # Find via drill/pad values
    via_map = {}
    for col_idx in range(4, len(row20)):
        if row20[col_idx]:
            via_val = str(row20[col_idx]).strip()
            if 'mil' in via_val.lower():
                via_map[col_idx + 1] = via_val
    
    # Combine track/spacing and via for each column
    for col_idx in range(5, min(len(row19), len(row20)) + 1):
        if col_idx in track_spacing_map or col_idx in via_map:
            column_specs[col_idx] = {
                'track_spacing': track_spacing_map.get(col_idx, ''),
                'via_drill_pad': via_map.get(col_idx, '')
            }
    
    # Extract rate data (starting from row 22)
    rates_data = []
    current_layers = None
    current_solder_mask = None
    
    for row_idx in range(22, ws.max_row + 1):
        row = [cell.value for cell in ws[row_idx]]
        
        # Skip empty rows
        if not any(row):
            continue
        
        # Extract specifications (columns 1-4)
        sl_no = row[0] if len(row) > 0 else None
        
        # Check if column 2 (layers) has a valid numeric value
        layers_value = row[1] if len(row) > 1 else None
        if layers_value is not None:
            if isinstance(layers_value, (int, float)):
                current_layers = int(layers_value)
            elif isinstance(layers_value, str):
                # Skip rows with text in layers column
                if any(keyword in layers_value.lower() for keyword in ['additional', 'cost', 'quantity', 'note', 'remark']):
                    continue
                try:
                    numbers = re.findall(r'\d+', layers_value)
                    if numbers:
                        current_layers = int(numbers[0])
                    else:
                        continue
                except:
                    continue
        
        # Check solder mask
        solder_mask_value = row[2] if len(row) > 2 else None
        if solder_mask_value is not None:
            if isinstance(solder_mask_value, str) and solder_mask_value.upper() in ['YES', 'NO']:
                current_solder_mask = solder_mask_value.upper()
        
        # Get thickness
        thickness = row[3] if len(row) > 3 and row[3] else None
        
        # Skip if we don't have essential data
        if not current_layers or not current_solder_mask or not thickness:
            continue
        
        # Skip if thickness contains keywords
        if isinstance(thickness, str):
            thickness_lower = thickness.lower()
            if any(keyword in thickness_lower for keyword in ['additional', 'cost', 'quantity', 'note', 'remark', 'total']):
                continue
        
        # Extract rates (columns 5+)
        for col_idx in range(4, len(row)):
            rate_value = row[col_idx]
            
            if isinstance(rate_value, (int, float)) and rate_value > 0:
                if col_idx + 1 in column_specs:
                    spec = column_specs[col_idx + 1]
                    rate_entry = {
                        'num_layers': current_layers,
                        'solder_mask': current_solder_mask,
                        'thickness': str(thickness).strip(),
                        'track_spacing': spec.get('track_spacing', ''),
                        'via_drill_pad': spec.get('via_drill_pad', ''),
                        'rate_per_sqcm': float(rate_value),
                        **basic_specs
                    }
                    rates_data.append(rate_entry)
    
    wb.close()
    
    return vendor_info, rates_data

def import_to_database(vendor_info, rates_data):
    """Import vendor and rate data into MySQL database"""
    connection = get_db_connection()
    cursor = connection.cursor()
    
    try:
        # Insert or update vendor
        cursor.execute("""
            INSERT INTO vendors (vendor_name, company_name, address, phone)
            VALUES (%s, %s, %s, %s)
            ON DUPLICATE KEY UPDATE
                company_name = VALUES(company_name),
                address = VALUES(address),
                phone = VALUES(phone)
        """, (vendor_info['name'], vendor_info['company_name'], 
              vendor_info['address'], vendor_info['phone']))
        
        vendor_id = cursor.lastrowid
        if not vendor_id:
            cursor.execute("SELECT vendor_id FROM vendors WHERE vendor_name = %s", 
                         (vendor_info['name'],))
            result = cursor.fetchone()
            vendor_id = result[0] if result else None
        
        if not vendor_id:
            print("Error: Could not get vendor ID")
            return
        
        print(f"Vendor ID: {vendor_id}")
        
        # Clear existing rates for this vendor
        cursor.execute("DELETE FROM vendor_rates WHERE vendor_id = %s", (vendor_id,))
        print(f"Cleared existing rates for vendor {vendor_info['name']}")
        
        # Insert rate data
        insert_query = """
            INSERT INTO vendor_rates (
                vendor_id, num_layers, solder_mask, thickness,
                track_spacing, via_drill_pad, rate_per_sqcm,
                material, surface_finish, copper_thickness,
                min_quantity, bbt_charges
            ) VALUES (
                %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s
            )
        """
        
        inserted_count = 0
        for rate in rates_data:
            try:
                cursor.execute(insert_query, (
                    vendor_id,
                    rate['num_layers'],
                    rate['solder_mask'],
                    rate['thickness'],
                    rate['track_spacing'],
                    rate['via_drill_pad'],
                    rate['rate_per_sqcm'],
                    rate.get('material'),
                    rate.get('surface_finish'),
                    rate.get('copper_thickness'),
                    rate.get('min_quantity'),
                    rate.get('bbt_charges')
                ))
                inserted_count += 1
            except Error as e:
                print(f"Error inserting rate: {e}")
                print(f"Rate data: {rate}")
                continue
        
        connection.commit()
        print(f"\nSuccessfully imported {inserted_count} rate records")
        
    except Error as e:
        print(f"Error importing data: {e}")
        connection.rollback()
    finally:
        cursor.close()
        connection.close()

def main():
    # Get Excel file path
    excel_path = os.path.join(os.path.dirname(__file__), '..', EXCEL_FILE_PATH)
    
    if not os.path.exists(excel_path):
        print(f"Error: Excel file not found at {excel_path}")
        print(f"Please place the Excel file at: {excel_path}")
        sys.exit(1)
    
    print("=" * 60)
    print(f"Importing Vendor: {VENDOR_NAME}")
    print("=" * 60)
    print(f"Excel file: {excel_path}\n")
    
    print("Parsing Excel file...")
    vendor_info, rates_data = parse_excel_data(excel_path, VENDOR_NAME)
    
    print(f"\nVendor: {vendor_info['company_name']}")
    print(f"Total rate records found: {len(rates_data)}")
    
    if rates_data:
        print("\nSample rate records:")
        for i, rate in enumerate(rates_data[:5], 1):
            print(f"{i}. Layers: {rate['num_layers']}, Solder Mask: {rate['solder_mask']}, "
                  f"Thickness: {rate['thickness']}, Track/Spacing: {rate['track_spacing']}, "
                  f"Via: {rate['via_drill_pad']}, Rate: {rate['rate_per_sqcm']}")
    
    print("\nImporting to database...")
    import_to_database(vendor_info, rates_data)
    print("\nImport completed!")

if __name__ == '__main__':
    main()