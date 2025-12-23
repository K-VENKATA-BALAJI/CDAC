from openpyxl import load_workbook
import os

excel_path = os.path.join(os.path.dirname(__file__), '..', 'HIQ_Rate contract.xlsx')
wb = load_workbook(excel_path)
ws = wb.active

print("=" * 80)
print("FULL STRUCTURE ANALYSIS")
print("=" * 80)

# Read rows 18-21 to understand the header structure
print("\nHeader Structure (Rows 18-21):")
print("-" * 80)

headers = {}
for row_idx in range(18, 22):
    row = [cell.value for cell in ws[row_idx]]
    print(f"\nRow {row_idx}:")
    for col_idx, val in enumerate(row, 1):
        if val:
            print(f"  Col {col_idx}: {val}")
            if row_idx == 18:  # Main header row
                headers[col_idx] = str(val)

# Now understand the rate columns - they seem to be combinations of track/spacing and via
print("\n\nUnderstanding Rate Column Structure:")
print("-" * 80)
print("Row 19 (Track/Spacing) and Row 20 (Via drill/pad) define the rate columns")
row19 = [cell.value for cell in ws[19]]
row20 = [cell.value for cell in ws[20]]

print("\nTrack/Spacing values (Row 19):")
for col_idx, val in enumerate(row19, 1):
    if val:
        print(f"  Col {col_idx}: {val}")

print("\nVia drill/pad values (Row 20):")
for col_idx, val in enumerate(row20, 1):
    if val:
        print(f"  Col {col_idx}: {val}")

# Read a few complete data rows to understand the pattern
print("\n\nSample Complete Data Rows:")
print("-" * 80)
for row_idx in range(34, min(50, ws.max_row + 1)):
    row = [cell.value for cell in ws[row_idx]]
    # Extract key fields
    sl_no = row[0] if len(row) > 0 else None
    layers = row[1] if len(row) > 1 else None
    solder_mask = row[2] if len(row) > 2 else None
    thickness = row[3] if len(row) > 3 else None
    
    # Get rates (columns 5 onwards)
    rates = []
    for col_idx in range(4, min(len(row), 20)):  # Check columns 5-20
        if isinstance(row[col_idx], (int, float)):
            rates.append((col_idx + 1, row[col_idx]))
    
    if sl_no or layers or rates:
        print(f"\nRow {row_idx}:")
        print(f"  Sl No: {sl_no}, Layers: {layers}, Solder Mask: {solder_mask}, Thickness: {thickness}")
        if rates:
            print(f"  Rates found in columns: {[f'Col {c}: {r}' for c, r in rates[:5]]}")

wb.close()


