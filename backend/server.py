import tornado.ioloop
import tornado.web
import tornado.httpserver
import json
import os
from openpyxl import load_workbook
from database import Database

class CORSHandler(tornado.web.RequestHandler):
    def set_default_headers(self):
        self.set_header("Access-Control-Allow-Origin", "*")
        self.set_header("Access-Control-Allow-Headers", "Content-Type")
        self.set_header("Access-Control-Allow-Methods", "GET, POST, PUT, DELETE, OPTIONS")

    def options(self):
        self.set_status(204)
        self.finish()

class GetSpecificationsHandler(CORSHandler):
    """Get PCB specification options from Excel file"""
    def get(self):
        try:
            excel_path = os.path.join(os.path.dirname(__file__), 'data', 'pcb_specifications.xlsx')
            wb = load_workbook(excel_path)
            ws = wb.active
            
            # Read specification options from Excel
            specifications = {}
            
            # Read from row 2 onwards (row 1 is header)
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0]:  # Field name exists
                    field_name = str(row[0]).strip()
                    options = [str(val).strip() for val in row[1:] if val is not None and str(val).strip()]
                    if options:
                        specifications[field_name] = options
            
            self.set_header("Content-Type", "application/json")
            self.write(json.dumps(specifications))
        except Exception as e:
            self.set_status(500)
            self.set_header("Content-Type", "application/json")
            self.write(json.dumps({"error": str(e)}))

class CalculateRatesHandler(CORSHandler):
    """Calculate rates based on PCB specifications"""
    def post(self):
        try:
            data = json.loads(self.request.body)
            specifications = data.get('specifications', {})
            
            excel_path = os.path.join(os.path.dirname(__file__), 'data', 'pcb_vendors.xlsx')
            wb = load_workbook(excel_path)
            ws = wb.active
            
            vendors = []
            # Read vendor data from Excel
            # Format: Vendor Name, Base Rate, Layer Multiplier, Material Multiplier, Size Multiplier
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0]:  # Vendor name exists
                    vendor_name = str(row[0])
                    base_rate = float(row[1]) if row[1] else 0
                    layer_mult = float(row[2]) if len(row) > 2 and row[2] else 1.0
                    material_mult = float(row[3]) if len(row) > 3 and row[3] else 1.0
                    size_mult = float(row[4]) if len(row) > 4 and row[4] else 1.0
                    
                    # Calculate rate based on specifications
                    num_layers_str = str(specifications.get('Num of Layers', '2'))
                    try:
                        num_layers = int(num_layers_str.split()[0])
                    except:
                        num_layers = 2
                    
                    # Apply multipliers (simplified calculation)
                    calculated_rate = base_rate * layer_mult * material_mult * size_mult
                    
                    # Add some variation based on specifications
                    if num_layers > 2:
                        calculated_rate *= (1 + (num_layers - 2) * 0.1)
                    
                    vendors.append({
                        "vendor_name": vendor_name,
                        "rate": round(calculated_rate, 2)
                    })
            
            # Sort by rate and assign rank
            vendors.sort(key=lambda x: x['rate'])
            for idx, vendor in enumerate(vendors, 1):
                vendor['rank'] = idx
                vendor['sno'] = idx
            
            self.set_header("Content-Type", "application/json")
            self.write(json.dumps(vendors))
        except Exception as e:
            self.set_status(500)
            self.set_header("Content-Type", "application/json")
            self.write(json.dumps({"error": str(e)}))

class SaveFileHandler(CORSHandler):
    """Save user file data"""
    def post(self):
        try:
            data = json.loads(self.request.body)
            email = data.get('email')
            file_name = data.get('file_name')
            file_description = data.get('file_description')
            specifications = data.get('specifications', {})
            vendors = data.get('vendors', [])
            
            db = Database()
            result = db.save_user_data(email, file_name, file_description, specifications, vendors)
            
            self.set_header("Content-Type", "application/json")
            self.write(json.dumps({"success": True, "message": "Data saved successfully"}))
        except Exception as e:
            self.set_status(500)
            self.set_header("Content-Type", "application/json")
            self.write(json.dumps({"error": str(e)}))

class LoadFileHandler(CORSHandler):
    """Load user file data by email"""
    def post(self):
        try:
            data = json.loads(self.request.body)
            email = data.get('email')
            
            db = Database()
            user_data = db.get_user_data(email)
            
            self.set_header("Content-Type", "application/json")
            if user_data:
                self.write(json.dumps({"success": True, "data": user_data}))
            else:
                self.set_status(404)
                self.write(json.dumps({"error": "No data found for this email"}))
        except Exception as e:
            self.set_status(500)
            self.set_header("Content-Type", "application/json")
            self.write(json.dumps({"error": str(e)}))

def make_app():
    return tornado.web.Application([
        (r"/api/specifications", GetSpecificationsHandler),
        (r"/api/calculate-rates", CalculateRatesHandler),
        (r"/api/save-file", SaveFileHandler),
        (r"/api/load-file", LoadFileHandler),
    ])

if __name__ == "__main__":
    app = make_app()
    server = tornado.httpserver.HTTPServer(app)
    server.listen(8888)
    print("Server running on http://localhost:8888")
    tornado.ioloop.IOLoop.current().start()

